import hashlib
import json
from pathlib import Path
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from jinja2 import Environment, FileSystemLoader
from docx import Document
from docx.shared import Inches

BASE_DIR = Path("/content") if Path("/content").exists() else Path.cwd()
FUENTE = BASE_DIR / "datos" / "operacion_centros.csv"
SALIDAS = BASE_DIR / "salidas"
SALIDAS.mkdir(exist_ok=True)

df = pd.read_csv(FUENTE, parse_dates=["fecha"])
df["deficit"] = (df["demanda_paquetes"] - df["capacidad_paquetes"]).clip(lower=0)

capacidad_total = int(df["capacidad_paquetes"].sum())
demanda_total = int(df["demanda_paquetes"].sum())
deficit_total = int(df["deficit"].sum())
cobertura = round(100 * (1 - deficit_total / demanda_total), 2)

def_centro = df.groupby("centro_id")["deficit"].sum()
centro_top = def_centro.idxmax()
deficit_top = int(def_centro.max())
pct_top = round((deficit_top / deficit_total) * 100, 1)
deficit_camp = round(df[df["campana_programada"] == 1]["deficit"].mean(), 1)
deficit_norm = round(df[df["campana_programada"] == 0]["deficit"].mean(), 1)

evol = df.groupby("fecha")[["capacidad_paquetes", "demanda_paquetes"]].sum()
fig, ax = plt.subplots(figsize=(9, 4))
ax.plot(evol.index, evol["demanda_paquetes"], label="Demanda", color="#c0392b", linewidth=1.5)
ax.plot(evol.index, evol["capacidad_paquetes"], label="Capacidad", color="#2c3e50", linewidth=1.5)
ax.set_title("Demanda vs Capacidad diaria de despacho (9 centros)")
ax.set_xlabel("Fecha")
ax.set_ylabel("Paquetes")
ax.legend()
fig.autofmt_xdate()
fig.tight_layout()
fig.savefig(SALIDAS / "grafico.png", dpi=150)
plt.close(fig)

pd.DataFrame({
    "indicador": ["capacidad_total_paquetes", "demanda_total_paquetes", "deficit_total_paquetes"],
    "valor": [capacidad_total, demanda_total, deficit_total]
}).to_csv(SALIDAS / "reporte_kpis.csv", index=False, encoding="utf-8")

sha = hashlib.sha256(FUENTE.read_bytes()).hexdigest()
with open(SALIDAS / "metadatos.json", "w", encoding="utf-8") as f:
    json.dump({
        "fuente": "operacion_centros.csv",
        "registros_entrada": int(len(df)),
        "sha256_fuente": sha
    }, f, ensure_ascii=False, indent=4)

ctx = {
    "titulo": "Informe semanal de capacidad de despacho - SmartRoute Logistics",
    "periodo": f"{df['fecha'].min().strftime('%Y-%m-%d')} a {df['fecha'].max().strftime('%Y-%m-%d')}",
    "capacidad_total": f"{capacidad_total:,}",
    "demanda_total": f"{demanda_total:,}",
    "deficit_total": f"{deficit_total:,}",
    "cobertura": f"{cobertura:.2f}",
    "hallazgo": f"El centro {centro_top} concentra {deficit_top:,} paquetes de déficit ({pct_top}% del total). En días con campaña programada, el déficit medio por registro asciende a {deficit_camp} paquetes frente a {deficit_norm} en días normales.",
    "recomendacion": f"Priorizar la ampliación de capacidad operativa o la redistribución de rutas hacia el centro {centro_top}, reforzando la logística de despacho especialmente durante las jornadas de campañas activas.",
    "limitacion": "Los datos analizados corresponden al periodo de operación y evalúan el déficit de forma independiente por jornada sin acumular rezago de paquetes no entregados para el día siguiente.",
    "registros": len(df),
    "sha256": sha
}

env = Environment(loader=FileSystemLoader(str(BASE_DIR)))
html_out = env.get_template("plantilla.html").render(**ctx)
(SALIDAS / "informe.html").write_text(html_out, encoding="utf-8")

doc = Document()
doc.add_heading(ctx["titulo"], level=0)
doc.add_paragraph(f"Periodo analizado: {ctx['periodo']}")
doc.add_heading("Indicadores clave", level=1)
tabla = doc.add_table(rows=1, cols=2)
tabla.style = "Light Grid Accent 1"
tabla.rows[0].cells[0].text = "Indicador"
tabla.rows[0].cells[1].text = "Valor"

for nom, val in [
    ("Capacidad total (paquetes)", ctx["capacidad_total"]),
    ("Demanda total (paquetes)", ctx["demanda_total"]),
    ("Déficit total (paquetes)", ctx["deficit_total"]),
    ("Cobertura (%)", ctx["cobertura"])
]:
    row = tabla.add_row().cells
    row[0].text = nom
    row[1].text = val

doc.add_heading("Gráfico de Comportamiento Operativo", level=1)
doc.add_picture(str(SALIDAS / "grafico.png"), width=Inches(6))
doc.add_heading("Hallazgo Principal", level=1)
doc.add_paragraph(ctx["hallazgo"])
doc.add_heading("Recomendación", level=1)
doc.add_paragraph(ctx["recomendacion"])
doc.add_heading("Limitación Real", level=1)
doc.add_paragraph(ctx["limitacion"])
doc.add_paragraph(f"Fuente: operacion_centros.csv | Registros: {ctx['registros']} | SHA-256: {sha}")
doc.save(SALIDAS / "informe.docx")
print("Reportes generados correctamente")
